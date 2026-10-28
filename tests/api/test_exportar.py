import sqlite3
from io import BytesIO

from openpyxl import load_workbook

URL = "/api/periodos/2026/11"


def libro(cliente):
    r = cliente.get(f"{URL}/exportar.xlsx")
    assert r.status_code == 200
    assert r.headers["content-disposition"] == 'attachment; filename="turnos-2026-11.xlsx"'
    return load_workbook(BytesIO(r.content))


def test_xlsx_del_calendario(cliente, equipo):
    cal = cliente.post(f"{URL}/generar", json={}).json()
    wb = libro(cliente)
    assert wb.sheetnames == ["Calendario", "Por médico", "Resumen"]

    hoja = wb["Calendario"]
    assert hoja["A1"].value == "Turnos — noviembre de 2026"
    assert hoja["A2"].value == "BORRADOR — calendario no publicado"
    assert hoja["C4"].value == "Mañana (08:00–15:00)"
    filas = {f[0].value.date().isoformat(): f for f in hoja.iter_rows(min_row=5, max_row=25) if f[0].value}
    assert len(filas) == 21
    assert filas["2026-11-02"][2].value == "Festivo: Traslado de Todos los Santos"
    # Ana tiene fija la mañana de los lunes.
    assert "Ana" in filas["2026-11-09"][2].value

    por_medico = wb["Por médico"]
    assert [c.value for c in por_medico[5]][0] == "Ana"
    assert "M" in [c.value for c in por_medico[5]]

    resumen = wb["Resumen"]
    totales = [fila[2].value for fila in resumen.iter_rows(min_row=5) if fila[0].value]
    assert sum(totales) == sum(m["asignado"] for m in cal["medicos"]) / 2 == 40


def test_xlsx_publicado_y_medio_turno(cliente, equipo):
    cliente.post(f"{URL}/generar", json={})
    cliente.post(f"{URL}/cambios", json=[
        {"medico_id": equipo[1], "fecha": "2026-11-12", "mitades": ["M1", "M2", "T1"]},
    ])
    cliente.post(f"{URL}/publicar")
    wb = libro(cliente)
    assert wb["Calendario"]["A2"].value.startswith("Publicado el ")
    tarde = next(f[3].value for f in wb["Calendario"].iter_rows(min_row=5) if f[0].value and f[0].value.day == 12)
    assert "Bea (½)" in tarde
    fila_bea = next(f for f in wb["Por médico"].iter_rows(min_row=5) if f[0].value == "Bea")
    assert "M+½T" in [c.value for c in fila_bea]


def test_xlsx_mes_vacio(cliente):
    wb = libro(cliente)
    assert wb["Calendario"].max_row >= 25


def test_copia_de_seguridad(cliente, equipo, tmp_path):
    r = cliente.get("/api/copia-de-seguridad")
    assert r.status_code == 200
    assert r.headers["content-disposition"].startswith('attachment; filename="turnos-copia-')
    fichero = tmp_path / "copia.sqlite"
    fichero.write_bytes(r.content)
    with sqlite3.connect(fichero) as conexion:
        assert conexion.execute("SELECT count(*) FROM medico").fetchone() == (4,)


def test_xlsx_con_urgencias_y_domicilios(cliente, equipo):
    cobertura = [
        {"dia_semana": d, "mitad": h, "minimo": 1 if h[0] == "M" else 2, "maximo": 1 if h[0] == "M" else 2}
        for d in range(5) for h in ("M1", "M2", "T1", "T2")
    ]
    cliente.put("/api/cobertura/semanal", json=cobertura)
    cliente.post(f"{URL}/generar", json={})
    wb = libro(cliente)
    tarde = next(f[3].value for f in wb["Calendario"].iter_rows(min_row=5) if f[0].value and f[0].value.day == 3)
    assert "(urgencias)" in tarde and "(domicilios)" in tarde
    resumen = wb["Resumen"]
    assert [c.value for c in resumen[4]][-2:] == ["Urgencias", "Domicilios"]
    assert sum(f[8].value for f in resumen.iter_rows(min_row=5) if f[0].value) == 20
    celdas = [c.value for fila in wb["Por médico"].iter_rows(min_row=5) for c in fila if isinstance(c.value, str)]
    assert any(v.endswith(" U") for v in celdas) and any(v.endswith(" D") for v in celdas)


def test_xlsx_de_urgencias_y_domicilios(cliente, equipo):
    cobertura = [
        {"dia_semana": d, "mitad": h, "minimo": 1 if h[0] == "M" else 2, "maximo": 1 if h[0] == "M" else 2}
        for d in range(5) for h in ("M1", "M2", "T1", "T2")
    ]
    cliente.put("/api/cobertura/semanal", json=cobertura)
    cal = cliente.post(f"{URL}/generar", json={}).json()
    nombres = {m["id"]: m["nombre"] for m in cal["medicos"]}
    r = cliente.get(f"{URL}/urgencias-domicilios.xlsx")
    assert r.status_code == 200
    assert r.headers["content-disposition"] == 'attachment; filename="urgencias-domicilios-2026-11.xlsx"'
    wb = load_workbook(BytesIO(r.content))
    assert wb.sheetnames == ["Urgencias y domicilios", "Por médico"]

    hoja = wb["Urgencias y domicilios"]
    assert hoja["A1"].value == "Urgencias y domicilios — noviembre de 2026"
    assert [c.value for c in hoja[4]] == ["Fecha", "Día", "Urgencias", "Domicilios"]
    assert hoja.page_setup.orientation == "portrait"
    filas = {f[0].value.date().isoformat(): [c.value for c in f[1:4]] for f in hoja.iter_rows(min_row=5) if f[0].value}
    assert filas["2026-11-02"][1].startswith("Festivo")  # traslado de Todos los Santos
    papeles = {(x["fecha"], x["rol"]): nombres[x["medico_id"]] for x in cal["roles"]}
    assert filas["2026-11-03"] == ["Martes", papeles["2026-11-03", "URGENCIAS"], papeles["2026-11-03", "DOMICILIOS"]]

    resumen = wb["Por médico"]
    assert [c.value for c in resumen[4]] == ["Médico", "Urgencias", "Domicilios", "Total"]
    totales = [f[1].value + f[2].value for f in resumen.iter_rows(min_row=5) if f[0].value]
    assert sum(totales) == 2 * 20 and totales == [f[3].value for f in resumen.iter_rows(min_row=5) if f[0].value]


def test_xlsx_de_urgencias_y_domicilios_sin_servicio(cliente, equipo):
    # Con 1 médico por la tarde no hay papeles: se indica en cada día.
    cliente.post(f"{URL}/generar", json={})
    r = cliente.get(f"{URL}/urgencias-domicilios.xlsx")
    hoja = load_workbook(BytesIO(r.content))["Urgencias y domicilios"]
    textos = [f[2].value for f in hoja.iter_rows(min_row=5) if f[0].value and f[0].value.day == 3]
    assert textos == ["Sin servicio (menos de dos médicos de tarde)"]


def test_xlsx_segun_la_presencia_sin_rangos(cliente, equipo):
    from io import BytesIO

    from openpyxl import load_workbook

    cliente.put("/api/ajustes", json={"modo_carga": "PRESENCIA"})
    cliente.post("/api/periodos/2026/11/generar", json={})
    libro = load_workbook(BytesIO(cliente.get("/api/periodos/2026/11/exportar.xlsx").content))
    resumen = libro["Resumen"] if "Resumen" in libro.sheetnames else libro.worksheets[-1]
    cabecera = [c.value for c in resumen[4]]
    assert "Mínimo" not in cabecera and "Situación" not in cabecera and "Urgencias" in cabecera
