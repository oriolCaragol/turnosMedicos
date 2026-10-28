"""Exportación del calendario de un mes a XLSX (diseño, 8).

Tres hojas:
- Calendario: un día por fila, con los médicos de mañana y de tarde.
- Por médico: un médico por fila y un día por columna (M, T, M+½T…), para
  que cada uno vea sus turnos de un vistazo.
- Resumen: turnos de cada médico frente a su rango del mes.

Y un informe aparte de urgencias y domicilios (`exportar_roles_xlsx`).
"""

from collections import defaultdict
from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from turnos.dominio.modelos import Mitad, Rol, Turno
from turnos.motor.resultado import TipoAviso
from turnos.servicios.periodos import Calendario

MESES = ("enero", "febrero", "marzo", "abril", "mayo", "junio", "julio",
         "agosto", "septiembre", "octubre", "noviembre", "diciembre")
DIAS = ("Lunes", "Martes", "Miércoles", "Jueves", "Viernes")
DIAS_CORTOS = ("L", "M", "X", "J", "V")

NEGRITA = Font(bold=True)
TITULO = Font(bold=True, size=14)
SUAVE = Font(italic=True, color="666666")
CABECERA = PatternFill("solid", fgColor="DBE8FF")
FESTIVO = PatternFill("solid", fgColor="ECEEF2")
FUERA = PatternFill("solid", fgColor="FDECEC")
FINO = Side(style="thin", color="C8CDD6")
BORDE = Border(left=FINO, right=FINO, top=FINO, bottom=FINO)
AJUSTE = Alignment(wrap_text=True, vertical="top")

JORNADAS = {
    frozenset({Mitad.M1, Mitad.M2}): "M",
    frozenset({Mitad.T1, Mitad.T2}): "T",
    frozenset({Mitad.M1, Mitad.M2, Mitad.T1}): "M+½T",
    frozenset({Mitad.M2, Mitad.T1, Mitad.T2}): "½M+T",
    frozenset(Mitad): "M+T",
}


def nombre_mes(anio: int, mes: int) -> str:
    return f"{MESES[mes - 1]} de {anio}"


def _turnos(medios: int) -> int | float:
    """Medios turnos → número de turnos (entero si es exacto, para que se vea "12" y no "12.0")."""
    return medios // 2 if medios % 2 == 0 else medios / 2


def _jornada(mitades: set[Mitad]) -> str:
    return JORNADAS.get(frozenset(mitades)) or "+".join(sorted(h.value for h in mitades))


def _cabecera(hoja: Worksheet, fila: int, textos: list[str]) -> None:
    for col, texto in enumerate(textos, start=1):
        celda = hoja.cell(fila, col, texto)
        celda.font, celda.fill, celda.border = NEGRITA, CABECERA, BORDE


def _titulo(hoja: Worksheet, cal: Calendario, texto: str) -> None:
    periodo = cal.datos.periodo
    hoja["A1"] = f"{texto} — {nombre_mes(cal.datos.anio, cal.datos.mes)}"
    hoja["A1"].font = TITULO
    if periodo is not None and periodo.publicado:
        hoja["A2"] = f"Publicado el {periodo.publicado_en:%d/%m/%Y}"
    else:
        hoja["A2"] = "BORRADOR — calendario no publicado"
    hoja["A2"].font = SUAVE


def _preparar_impresion(hoja: Worksheet, orientacion: str = "landscape") -> None:
    hoja.page_setup.orientation = orientacion
    hoja.page_setup.paperSize = hoja.PAPERSIZE_A4
    hoja.page_setup.fitToWidth = 1
    hoja.page_setup.fitToHeight = 0
    hoja.sheet_properties.pageSetUpPr.fitToPage = True


def exportar_xlsx(cal: Calendario, horarios: dict[Turno, str]) -> bytes:
    """`horarios`: texto del horario de cada turno, p. ej. "08:00–15:00"."""
    nombres = {mm.medico.id: mm.medico.nombre for mm in cal.datos.medicos}
    mitades: dict[tuple[int, date], set[Mitad]] = defaultdict(set)
    for a in cal.asignaciones:
        mitades[a.medico_id, a.fecha].add(a.mitad)
    roles = {(r.medico_id, r.fecha): r.rol for r in cal.roles}

    libro = Workbook()
    _hoja_calendario(libro.active, cal, nombres, mitades, horarios, roles)
    _hoja_por_medico(libro.create_sheet("Por médico"), cal, mitades, roles)
    _hoja_resumen(libro.create_sheet("Resumen"), cal)
    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()


def _hoja_calendario(hoja, cal, nombres, mitades, horarios, roles) -> None:
    hoja.title = "Calendario"
    _titulo(hoja, cal, "Turnos")
    _cabecera(hoja, 4, ["Fecha", "Día", f"Mañana ({horarios[Turno.MANANA]})", f"Tarde ({horarios[Turno.TARDE]})"])
    fila = 5
    for dia in cal.datos.dias_semana:
        hoja.cell(fila, 1, dia).number_format = "DD/MM/YYYY"
        hoja.cell(fila, 2, DIAS[dia.weekday()])
        festivo = cal.datos.festivos.get(dia)
        if festivo:
            hoja.cell(fila, 3, f"Festivo: {festivo}").font = SUAVE
            hoja.merge_cells(start_row=fila, start_column=3, end_row=fila, end_column=4)
            for col in range(1, 5):
                hoja.cell(fila, col).fill = FESTIVO
        else:
            for col, turno in ((3, Turno.MANANA), (4, Turno.TARDE)):
                textos = []
                for medico_id, nombre in sorted(nombres.items(), key=lambda i: i[1]):
                    propias = {h for h in mitades.get((medico_id, dia), ()) if h.turno is turno}
                    if propias:
                        rol = roles.get((medico_id, dia)) if turno is Turno.TARDE else None
                        textos.append(
                            nombre + (" (½)" if len(propias) == 1 else "") + (f" ({rol.nombre})" if rol else "")
                        )
                hoja.cell(fila, col, ", ".join(textos)).alignment = AJUSTE
        for col in range(1, 5):
            hoja.cell(fila, col).border = BORDE
        fila += 1
    incumplimientos = [a for a in cal.avisos if a.tipo is TipoAviso.INCUMPLIMIENTO]
    fila += 1
    hoja.cell(fila, 1, "½ = medio turno (mitad del turno contigua a la jornada).").font = SUAVE
    if incumplimientos:
        fila += 2
        hoja.cell(fila, 1, "Reglas que no se cumplen:").font = NEGRITA
        for aviso in incumplimientos:
            fila += 1
            hoja.cell(fila, 1, aviso.mensaje)
    for col, ancho in zip("ABCD", (12, 11, 60, 60)):
        hoja.column_dimensions[col].width = ancho
    hoja.freeze_panes = "A5"
    _preparar_impresion(hoja)


def _hoja_por_medico(hoja, cal, mitades, roles) -> None:
    _titulo(hoja, cal, "Turnos por médico")
    dias = cal.datos.dias_semana
    _cabecera(hoja, 4, ["Médico"] + [f"{DIAS_CORTOS[d.weekday()]} {d.day}" for d in dias])
    fila = 5
    for mm in cal.datos.medicos:
        hoja.cell(fila, 1, mm.medico.nombre).border = BORDE
        for col, dia in enumerate(dias, start=2):
            celda = hoja.cell(fila, col)
            celda.border = BORDE
            celda.alignment = Alignment(horizontal="center")
            if dia in cal.datos.festivos:
                celda.fill = FESTIVO
            elif propias := mitades.get((mm.medico.id, dia)):
                rol = roles.get((mm.medico.id, dia))
                celda.value = _jornada(propias) + (" " + ("U" if rol is Rol.URGENCIAS else "D") if rol else "")
        fila += 1
    fila += 1
    hoja.cell(
        fila, 1, "M = mañana · T = tarde · ½ = medio turno · U = urgencias · D = domicilios · gris = festivo"
    ).font = SUAVE
    hoja.column_dimensions["A"].width = 24
    for col in range(2, len(dias) + 2):
        hoja.column_dimensions[get_column_letter(col)].width = 6.5
    hoja.freeze_panes = "B5"
    _preparar_impresion(hoja)


def _hoja_resumen(hoja, cal) -> None:
    """Con el modo de carga «según la presencia» no hay rangos: sin mínimo, máximo ni situación."""
    rangos = not cal.datos.por_presencia
    _titulo(hoja, cal, "Resumen de carga")
    _cabecera(hoja, 4, [
        "Médico", "Código", "Turnos", "Mañanas", "Tardes",
        *(["Mínimo", "Máximo", "Situación"] if rangos else []), "Urgencias", "Domicilios",
    ])
    fila = 5
    for mm in cal.datos.medicos:
        carga = cal.carga.get(mm.medico.id, {})
        total = sum(carga.values())
        valores = [
            mm.medico.nombre,
            mm.medico.codigo,
            _turnos(total),
            _turnos(carga.get(Turno.MANANA, 0)),
            _turnos(carga.get(Turno.TARDE, 0)),
        ]
        if not rangos:
            pass
        elif mm.rango is None:
            valores += ["", "", "Sin carga definida"]
        else:
            r = mm.rango.final
            situacion = (
                "Por debajo del mínimo" if total < r.minimo
                else "Por encima del máximo" if total > r.maximo
                else "Dentro del rango"
            )
            valores += [_turnos(r.minimo), _turnos(r.maximo), situacion]
        fuera = str(valores[-1]).startswith("Por ")
        veces = cal.veces_rol.get(mm.medico.id, {})
        valores += [veces.get(Rol.URGENCIAS, 0), veces.get(Rol.DOMICILIOS, 0)]
        for col, valor in enumerate(valores, start=1):
            celda = hoja.cell(fila, col, valor)
            celda.border = BORDE
            if fuera:
                celda.fill = FUERA
        fila += 1
    anchos = (24, 9, 9, 9, 9, 9, 9, 22, 11, 11) if rangos else (24, 9, 9, 9, 9, 11, 11)
    for col, ancho in zip("ABCDEFGHIJ", anchos):
        hoja.column_dimensions[col].width = ancho
    _preparar_impresion(hoja)


# --- Informe de urgencias y domicilios --------------------------------------------------


def exportar_roles_xlsx(cal: Calendario) -> bytes:
    """Urgencias y domicilios del mes: quién hace cada papel cada día, y cuántas veces
    cada médico."""
    libro = Workbook()
    _hoja_roles(libro.active, cal)
    _hoja_resumen_roles(libro.create_sheet("Por médico"), cal)
    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()


def _hoja_roles(hoja, cal: Calendario) -> None:
    hoja.title = "Urgencias y domicilios"
    _titulo(hoja, cal, "Urgencias y domicilios")
    nombres = {mm.medico.id: mm.medico.nombre for mm in cal.datos.medicos}
    papel = {(r.fecha, r.rol): nombres.get(r.medico_id, "") for r in cal.roles}
    _cabecera(hoja, 4, ["Fecha", "Día", "Urgencias", "Domicilios"])
    fila = 5
    for dia in cal.datos.dias_semana:
        hoja.cell(fila, 1, dia).number_format = "DD/MM/YYYY"
        hoja.cell(fila, 2, DIAS[dia.weekday()])
        festivo = cal.datos.festivos.get(dia)
        if festivo:
            hoja.cell(fila, 3, f"Festivo: {festivo}").font = SUAVE
            hoja.merge_cells(start_row=fila, start_column=3, end_row=fila, end_column=4)
            for col in range(1, 5):
                hoja.cell(fila, col).fill = FESTIVO
        elif not cal.datos.problema.requiere_roles(dia):
            hoja.cell(fila, 3, "Sin servicio (menos de dos médicos de tarde)").font = SUAVE
            hoja.merge_cells(start_row=fila, start_column=3, end_row=fila, end_column=4)
        else:
            for col, rol in ((3, Rol.URGENCIAS), (4, Rol.DOMICILIOS)):
                nombre = papel.get((dia, rol))
                celda = hoja.cell(fila, col, nombre or "Sin asignar")
                if not nombre:
                    celda.font, celda.fill = SUAVE, FUERA
        for col in range(1, 5):
            hoja.cell(fila, col).border = BORDE
        fila += 1
    for col, ancho in zip("ABCD", (12, 11, 32, 32)):
        hoja.column_dimensions[col].width = ancho
    hoja.freeze_panes = "A5"
    _preparar_impresion(hoja, "portrait")


def _hoja_resumen_roles(hoja, cal: Calendario) -> None:
    _titulo(hoja, cal, "Urgencias y domicilios por médico")
    _cabecera(hoja, 4, ["Médico", "Urgencias", "Domicilios", "Total"])
    fila = 5
    for mm in sorted(cal.datos.medicos, key=lambda m: m.medico.nombre):
        veces = cal.veces_rol.get(mm.medico.id, {})
        urgencias, domicilios = veces.get(Rol.URGENCIAS, 0), veces.get(Rol.DOMICILIOS, 0)
        if not (urgencias or domicilios):
            continue
        for col, valor in enumerate([mm.medico.nombre, urgencias, domicilios, urgencias + domicilios], start=1):
            hoja.cell(fila, col, valor).border = BORDE
        fila += 1
    for col, ancho in zip("ABCD", (28, 11, 11, 9)):
        hoja.column_dimensions[col].width = ancho
    _preparar_impresion(hoja, "portrait")

